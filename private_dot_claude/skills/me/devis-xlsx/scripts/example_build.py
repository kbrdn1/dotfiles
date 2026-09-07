"""
Exemple complet : construit un classeur de devis de bout en bout avec devis_xlsx.

Montre les 4 briques : params_block, estimation_sheet (back/front + sous-totaux),
une Synthèse câblée par formules sur les sous-totaux, et notes_sheet.
Se termine par un recalc qui doit afficher 0 erreur.

Lancer :  python scripts/example_build.py
"""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import Workbook
import devis_xlsx as dx

wb = Workbook()
syn = wb.active
syn.title = "Synthese"
est = wb.create_sheet("Estimation")
notes = wb.create_sheet("Notes")

# ---------- Synthèse : titre + paramètres ----------
dx.title(syn, "Devis — Exemple de prestation",
         "Démo devis_xlsx · jours-homme · HT sauf mention · TJM à saisir")
for col, w in {"A": 34, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12}.items():
    syn.column_dimensions[col].width = w

params = [
    {"key": "tjm", "label": "TJM (€/jour)", "value": 500, "fmt": dx.EUR,
     "note": "Exemple — remplacer par ton TJM réel"},
    {"key": "tva", "label": "TVA (%)", "value": 0.20, "fmt": dx.PCT, "note": "0 si non assujetti"},
    {"key": "buf", "label": "Buffer de risque (%)", "value": 0.15, "fmt": dx.PCT,
     "note": "Justifié par les zones grises (voir Notes)"},
    {"key": "tog_opt", "label": "Inclure le lot optionnel ? (1/0)", "value": 1, "fmt": dx.INT,
     "note": "Sélecteur de périmètre"},
]
refs, _ = dx.params_block(syn, params)   # refs['tjm'] == 'B5', etc.
TJM = f"Synthese!${refs['tjm']}"
TVA = f"${refs['tva']}"
BUF = f"${refs['buf']}"
TOG = f"${refs['tog_opt']}"

# ---------- Estimation : back + front + lot optionnel ----------
blocks = [
    {"section": "BACK (Laravel)", "subtotal": "Sous-total Back",
     "lines": [
         {"cat": "Back", "poste": "Module CRUD entité X", "cls": "Extension", "low": 2, "high": 4,
          "files": "CRUDController, Repository, Resource"},
         {"cat": "Back", "poste": "Workflow à états", "cls": "Greenfield (gabarit)", "low": 3, "high": 6,
          "files": "Service de transitions existant"},
     ]},
    {"section": "FRONT (Nuxt)", "subtotal": "Sous-total Front",
     "lines": [
         {"cat": "Front", "poste": "Page admin config-driven", "cls": "Extension", "low": 2, "high": 3,
          "files": "useAdminTablePage, store + config"},
         {"cat": "Front", "poste": "Formulaire Zod + VeeValidate", "cls": "Extension", "low": 1.5, "high": 3,
          "files": "schemas/*.ts"},
     ]},
    {"section": "LOT OPTIONNEL", "subtotal": "Sous-total Option",
     "lines": [
         {"cat": "Option", "poste": "Tableau de bord stats", "cls": "Greenfield partiel", "low": 2, "high": 4,
          "files": "charts existants"},
     ]},
]
subs = dx.estimation_sheet(est, blocks, tjm_ref=TJM)
# subs == {"Sous-total Back": r1, "Sous-total Front": r2, "Sous-total Option": r3}
BACK = subs["Sous-total Back"]
FRONT = subs["Sous-total Front"]
OPT = subs["Sous-total Option"]

# ---------- Synthèse : récap câblé par formules ----------
hdr = 11
heads = ["Bloc", "Jours bas", "Jours haut", "€ bas", "€ médian", "€ haut"]
for i, h in enumerate(heads):
    dx.style(syn.cell(row=hdr, column=1 + i, value=h), bold=True, fill="light",
             align=("left" if i == 0 else "center"), border=True)

# Dev = Back + Front (toujours) ; Option pilotée par le sélecteur
rows = [
    ("Back", f"=Estimation!D{BACK}", f"=Estimation!E{BACK}", "link"),
    ("Front", f"=Estimation!D{FRONT}", f"=Estimation!E{FRONT}", "link"),
    ("Lot optionnel (× sélecteur)", f"=Estimation!D{OPT}*{TOG}", f"=Estimation!E{OPT}*{TOG}", "toggle"),
]
r = hdr + 1
first = r
for label, fb, fh, kind in rows:
    dx.style(syn.cell(row=r, column=1, value=label), border=True)
    syn.cell(row=r, column=2, value=fb)
    syn.cell(row=r, column=3, value=fh)
    col_color = "green" if kind == "link" else "black"
    dx.style(syn.cell(row=r, column=2), color=col_color, numfmt=dx.DAY, align="center", border=True)
    dx.style(syn.cell(row=r, column=3), color=col_color, numfmt=dx.DAY, align="center", border=True)
    # € = jours * TJM ; médian = moyenne
    syn.cell(row=r, column=4, value=f"=B{r}*{TJM}")
    syn.cell(row=r, column=6, value=f"=C{r}*{TJM}")
    syn.cell(row=r, column=5, value=f"=(D{r}+F{r})/2")
    for c in (4, 5, 6):
        dx.style(syn.cell(row=r, column=c), numfmt=dx.EUR, align="center", border=True)
    r += 1
last = r - 1

# Sous-total dev pur
dx.style(syn.cell(row=r, column=1, value="Dev pur"), bold=True, border=True, fill="light")
for c, L in zip((2, 3, 4, 5, 6), ("B", "C", "D", "E", "F")):
    syn.cell(row=r, column=c, value=f"=SUM({L}{first}:{L}{last})")
    dx.style(syn.cell(row=r, column=c), bold=True, numfmt=(dx.DAY if c in (2, 3) else dx.EUR),
             align="center", border=True, fill="light")
devpur = r
r += 1

# Total avec buffer
dx.style(syn.cell(row=r, column=1, value="TOTAL projet (HT, buffer inclus)"), bold=True, color="white", fill="navy", border=True)
for c, L in zip((2, 3, 4, 5, 6), ("B", "C", "D", "E", "F")):
    syn.cell(row=r, column=c, value=f"={L}{devpur}*(1+{BUF})")
    dx.style(syn.cell(row=r, column=c), bold=True, color="white", fill="navy",
             numfmt=(dx.DAY if c in (2, 3) else dx.EUR), align="center", border=True)
total = r
r += 1

# TTC (sur médian)
dx.style(syn.cell(row=r, column=1, value="dont TTC (médian)"), border=True)
syn.cell(row=r, column=5, value=f"=E{total}*(1+{TVA})")
dx.style(syn.cell(row=r, column=5), bold=True, numfmt=dx.EUR, align="center", border=True, fill="light")

# ---------- Notes ----------
dx.notes_sheet(notes, [
    {"heading_top": "Écarts, hypothèses & décisions",
     "title": "Écarts source ↔ code", "numbered": True,
     "items": [
         "Exemple : un sous-total de l'estimation source ne correspondait pas à la somme de ses lignes (lot oublié). Le classeur recalcule par formule et documente l'écart ici plutôt que de recopier un total faux.",
     ]},
    {"title": "Hypothèses", "items": [
        "1 dev mid/senior connaissant le repo (pas d'onboarding).",
        "Tests unitaires de base inclus dans le dev ; recette transverse à part si besoin.",
        "Jours uniquement ; TJM, TVA, buffer et périmètre sont des paramètres à saisir dans l'onglet Synthese.",
    ]},
])

dx.move_first(wb, "Synthese")
out = "/tmp/exemple_devis.xlsx"
wb.save(out)
print("Classeur écrit :", out)

res = dx.recalc(out)
print("Recalc :", res)
